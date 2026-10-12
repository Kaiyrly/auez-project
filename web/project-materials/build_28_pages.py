"""Qalam: 28 pages with typography taken from the user-supplied sample."""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from package_content import *
HERE=Path(__file__).parent;ROOT=HERE.parent;OUT=ROOT/'downloads'
annotation=(HERE/'annotation-190-words.txt').read_text().strip()
assert len(annotation.split())==190
material=deepcopy(pages)
material[0]=(material[0][0],[])
material[2]=('АННОТАЦИЯ',[P(annotation),P('Кілт сөздер: мектеп кітапханасы, Qalam, электрондық каталог, брондау, кітап қоры, оқырман, сауалнама.')])
material[5]=('I. МЕКТЕП КІТАПХАНАСЫН ЦИФРЛАНДЫРУДЫҢ ТЕОРИЯЛЫҚ НЕГІЗДЕРІ',[H('1.1 Электрондық каталог және ақпарат іздеу')]+material[5][1])
material[6]=('1.2 Кітап қоры және дана есебі',material[6][1])
material[7]=('1.3 Оқырман қызығушылығы мен кері байланыс',material[7][1])
material[8]=('II. QALAM ВЕБ-ЖҮЙЕСІН ЖОБАЛАУ',[H('2.1 Жүйенің құрылымы және интерфейс логикасы')]+material[8][1])
material[11]=('2.2 Брондау және кітап беру есебі',material[11][1])
material[14]=('2.3 Кітап ұсыну және оқу мақсаты',material[14][1])
material[15]=('III. НӘТИЖЕЛЕР, ТАЛҚЫЛАУ ЖӘНЕ ПРАКТИКАЛЫҚ МАҢЫЗЫ',[H('3.1 Зерттеу және тексеру әдісі')]+material[15][1])
material[18]=('3.2 Бағдарламалық тексеру нәтижелері',material[18][1])
material[20]=('3.3 Талқылау және ғылыми қорытынды',material[20][1])
material[22][1][:]=material[22][1][:4]+[P(RULE_NOTE)]
material[24][1].append(P('2026 жылғы «Зерде» конкурсын ұйымдастыру және өткізу ережелері. Пайдаланушы ұсынған «ЗЕРДЕ ЕРЕЖЕ 2026.pdf»: 4.5–4.13-тармақтар және 4–8-қосымшалар.'))
material[23]=('ҚОРЫТЫНДЫ ЖӘНЕ РЕФЛЕКСИЯ',material[23][1])
material[24]=('ПАЙДАЛАНЫЛҒАН ДЕРЕККӨЗДЕР',material[24][1])
material.extend([
('ҚОСЫМША А. БАСТАПҚЫ САУАЛНАМА',[
 P('Мақсаты — сайтты жасауға дейін оқушылардың қажеттілігін анықтау. Аты-жөнін жазу міндетті емес. Қатысу ерікті.'),
 H('1. Керекті кітапты табу саған оңай ма?'),P('□ Оңай   □ Кейде қиын   □ Қиын'),
 H('2. Кітап таңдағанда не қиын?'),P('Бірнеше жауап: □ Не оқитынымды білмеймін. □ Қандай кітап барын білмеймін. □ Іздеуге көп уақыт кетеді. □ Қиындық жоқ.'),
 H('3. Кітапхана сайты болса, қолданар ма едің?'),P('□ Иә   □ Мүмкін   □ Жоқ'),
 H('4. Қандай мүмкіндіктер қажет? Үшеуін таңда.'),P('□ Іздеу □ Бос дананы көру □ Брондау □ Кітап сипаттамасы □ Кітап ұсыну □ Пікір мен баға □ Кітап тарихы'),
 H('5. Кітап оқуға қызығушылығың қандай?'),P('1 — өте төмен, 5 — өте жоғары: □ 1 □ 2 □ 3 □ 4 □ 5'),
 H('6. Сайтты қай құрылғыдан ашар едің?'),P('□ Телефон □ Компьютер / ноутбук □ Планшет'),
 H('7. Болашақ сайтқа қандай ұсынысың бар?'),P('________________________________________________'),
 P('Бланк зерттеу жүргізуге арналған. Оқушы жауаптары ретінде берілмеген.')]),
('ҚОСЫМША Ә. КЕЙІНГІ САУАЛНАМА',[
 P('Мақсаты — сайтты қолдану тәжірибесін білу. Жауапты өз тәжірибеңе сүйеніп бер. 1 — мүлде келіспеймін, 5 — толық келісемін.'),
 T(['Сұрақ','Баға'],[['1. Сайтты қолдану оңай болды.','1 / 2 / 3 / 4 / 5'],['2. Керекті кітапты тез таптым.','1 / 2 / 3 / 4 / 5'],['3. Бос дананы көру ыңғайлы.','1 / 2 / 3 / 4 / 5'],['4. Онлайн брондау маған пайдалы.','1 / 2 / 3 / 4 / 5']]),
 H('5. Кітап оқуға қызығушылығың қандай?'),P('1 — өте төмен, 5 — өте жоғары: □ 1 □ 2 □ 3 □ 4 □ 5'),
 H('6. Қай мүмкіндік көбірек ұнады?'),P('□ Іздеу □ Брондау □ Кітап ұсыну □ Пікір мен баға □ Рейтинг □ Оқу мақсаты'),
 H('7. Сайтты қайта қолданғың келе ме?'),P('□ Иә   □ Мүмкін   □ Жоқ'),
 H('8. Нені өзгерту немесе қосу керек?'),P('________________________________________________'),
 P('Екі сауалнамадағы 5-сұрақ бірдей. Жауаптар аты-жөнсіз кодпен сәйкестендіріледі. Қорытынды нақты жауаптар жиналған соң жасалады.')]),
('ҚОСЫМША Б. СЫНАҚ ХАТТАМАСЫ',[
 P('Бұл кесте қағаз тізім мен сайттағы кітап іздеуді салыстыруға арналған. Әзірге толтырылмаған. Қатысушының аты-жөні орнына код жазылады.'),
 P('Күні: __________  Қатысушы коды: __________\nҚұрылғы: __________  Тапсырма жиыны: __________'),
 T(['Тапсырма','Қағаз: секунд','Сайт: секунд','Дұрыс жауап'],[['Кітапты атауымен табу','____','____','____'],['Автордың кітабын табу','____','____','____'],['Жанр бойынша таңдау','____','____','____'],['Бос данасын анықтау','____','____','____']]),
 H('Сынақ тәртібі'),P('Қатысушылардың бір бөлігі алдымен қағазбен, екінші бөлігі алдымен сайтпен жұмыс істейді. Екі тәсілге күрделілігі бірдей тапсырмалар беріледі. Уақыт тапсырма оқылғаннан жауап табылғанға дейін өлшенеді.'),
 H('Нәтижені есептеу'),P('Орташа уақыт — барлық уақыттың қосындысын қатысушы санына бөлу. Дұрыс жауап үлесі — дұрыс жауап санын барлық жауап санына бөліп, 100-ге көбейту.'),
 H('Бақылау жазбасы'),P('Кездескен қиындық: ______________________________\nОқушы ұсынысы: _________________________________'),
 P('Сынаққа дейін мектептің қажет рұқсаттары алынады. Нақты уақыт пен жауап алынбайынша, тиімділік пайызы жазылмайды.')])])
assert len(material)==28
sample=Document(HERE/'Qalam-format-template.docx');proto=[deepcopy(p._p) for p in sample.paragraphs]
doc=sample
for e in list(doc._element.body):
 if not e.tag.endswith('}sectPr'):doc._element.body.remove(e)
def para(value,kind=6,size=None,bold=False,italic=False):
 p=doc.add_paragraph();pr=proto[kind].find(qn('w:pPr'))
 if pr is not None:p._p.insert(0,deepcopy(pr))
 # Keep the sample's pPr, but never carry a pagination flag into ordinary text.
 if kind not in [4,7]:p.paragraph_format.page_break_before=False
 r=p.add_run(value);r.font.name='Times New Roman'
 if size:r.font.size=Pt(size)
 if bold:r.bold=True
 if italic:r.italic=True
 return p

table_number=0
def table(headers,rows):
 global table_number
 table_number+=1
 para(f'{table_number}-кесте. '+headers[0],9,size=11)
 t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for i,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=i==0

 para('Дереккөз: Qalam жүйесінің тексерілген сипаттамасы немесе зерттеуге арналған толтырылатын бланк.',10,size=10,italic=True)

def block(item):
 kind,v=item
 if kind=='p':para(v)
 elif kind=='h':para(v,8,bold=True)
 elif kind=='bullets':
  for x in v:para('• '+x)
 elif kind=='table':table(*v)
 elif kind=='image':
  name,caption,width=v;path=HERE/'assets'/name;im=Image.open(path);width=min(width,7.5*im.width/im.height)
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Cm(width))
  para(caption,9,size=11)

for i,(title,content) in enumerate(material):
 if i:doc.add_page_break()
 if i==0:
  for s in ['Абай облысы білім басқармасы ММ','«Дарындылық пен қосымша білім беруді дамытудың облыстық ғылыми-әдістемелік «Дарын» орталығы» КМҚК','2–8-сынып оқушылары арасында өткізілетін «Зерде» республикалық зерттеу жобалары мен шығармашылық жұмыстары конкурсы']:para(s,0,bold=True)
  for s in ['Тіркеу коды (шифр): ______________________________','Бағыт: жаратылыстану-математикалық','Секция: Информатика, жасанды интеллект, робототехника және STEM-инженерия','Жас санаты: Орта буын (5–8 сынып)']:para(s,1)
  para('ҒЫЛЫМИ ЖОБА',2,size=16,bold=True);para('«'+TITLE+'»',3,size=16,bold=True)
  for s in ['Қатысушылар: '+AUTHORS,'Сыныбы: 7-сынып','Пәні: информатика','Ғылыми жетекші: '+LEADER,'Білім беру ұйымы: '+SCHOOL,'Өңір: Абай облысы']:para(s,1)
  para(YEAR,0,bold=True)
 elif i==1:
  para(title,4,bold=True)
  toc=[('Аннотация',3),('Кіріспе',4),('Зерттеудің мақсаты мен жоспары',5)]
  toc += [(material[n-1][0].capitalize() if n not in [6,9,16] else material[n-1][0],n) for n in [6,7,8,9,12,15,16,19,21]]
  toc += [('Этика және деректерді сақтау',22),('Жасанды интеллект және авторлық үлес',23),('Қорытынды және рефлексия',24),('Пайдаланылған дереккөздер',25),('Қосымша А. Бастапқы сауалнама',26),('Қосымша Ә. Кейінгі сауалнама',27),('Қосымша Б. Сынақ хаттамасы',28)]
  for label,n in toc:para(label+'\t'+str(n),5,size=13,bold=n in [6,9,16])
 else:
  para(title,7 if i in [5,8,15] else 4,bold=True)
  for item in content:block(item)
doc.core_properties.keywords='Qalam, мектеп кітапханасы, каталог, брондау, Зерде';doc.core_properties.title=TITLE;doc.core_properties.author=AUTHORS;doc.core_properties.last_modified_by='Qalam';doc.core_properties.subject='28 беттік ғылыми жоба; аннотация 190 сөз';doc.core_properties.comments='Үлгі бойынша Qalam жобасына бейімделген.'
# Rule 4.5: no visible page number on the cover.
section=doc.sections[0];section.different_first_page_header_footer=True
for p in section.first_page_footer.paragraphs:p.clear()
doc.save(OUT/'Qalam-scientific-project-28-pages.docx')
# Rule 4.8: judging copy must not identify pupils, supervisor or school.
anon=Document(OUT/'Qalam-scientific-project-28-pages.docx')
for p in list(anon.paragraphs):
 if p.text.startswith(('Қатысушылар:','Ғылыми жетекші:','Білім беру ұйымы:','Өңір:')):
  p._element.getparent().remove(p._element)
for p in anon.paragraphs:
 for r in p.runs:
  updated=r.text.replace(SCHOOL,'мектеп').replace(AUTHORS,'1-автор және 2-автор').replace(LEADER,'ғылыми жетекші').replace('«Қаныш Сәтпаев атындағы Жәнтікей орта мектебі»','мектеп').replace('github.com/Kaiyrly/auez-project','Жоба коды әкімшілік құжаттар топтамасында көрсетілген')
  if updated!=r.text:r.text=updated
for t in anon.tables:
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    for r in p.runs:
     old=r.text;new=old.replace('Айбын','1-автор').replace('Елжан','2-автор').replace(AUTHORS,'1-автор және 2-автор').replace(LEADER,'ғылыми жетекші').replace(SCHOOL,'мектеп')
     if new!=old:r.text=new
for rel in anon.part.rels.values():
 if rel.reltype.endswith('/image'):
  # The poster contains school and author identifiers; replace this visual only.
  from PIL import Image
  import io
  im=Image.open(io.BytesIO(rel.target_part.blob))
  if im.size==Image.open(HERE/'assets'/'poster.png').size:
   rel.target_part._blob=(HERE/'assets'/'architecture.png').read_bytes()
for p in anon.paragraphs:
 if p.text.startswith('4-сурет.'):
  for r in p.runs:r.text='4-сурет. Жүйенің ақпарат ағыны. Дереккөз: Qalam құрылымының сызбасы.'
anon.core_properties.author='';anon.core_properties.last_modified_by='';anon.core_properties.comments='Қазыларға арналған нұсқа. Тіркеу кодын толтыру қажет.'
anon.save(OUT/'Qalam-scientific-project-28-pages-anonymous.docx')
print('28 planned pages; annotation:',len(annotation.split()),'words')

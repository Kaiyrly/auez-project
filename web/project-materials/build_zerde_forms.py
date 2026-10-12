"""Prepare Qalam forms using appendices 3–7 of the supplied Zerde 2026 rules."""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from package_content import *
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'downloads/Qalam-project-package'
def create(title,annex):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=Cm(3);s.right_margin=s.top_margin=s.bottom_margin=Cm(2)
 st=d.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(14);st.paragraph_format.line_spacing=1.5;st.paragraph_format.space_after=Pt(5)
 d.core_properties.author=AUTHORS;d.core_properties.title=title
 p=d.add_paragraph(annex+'-қосымша' if annex.isdigit() else annex);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(title).bold=True
 return d

def common(d,section=True):
 for s in ['Тіркеу коды (шифры): ____________________','Бағыты: жаратылыстану-математикалық']+(['Секция атауы: '+SECTION] if section else [])+['Жоба тақырыбы: «'+TITLE+'»']:
  d.add_paragraph(s)
def grid(d,headers,rows,size=11):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for i,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3)
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=i==0
 return t

def signatures(d,responsible=False):
 d.add_paragraph('1-автордың қолы: ____________________\n2-автордың қолы: ____________________\nҒылыми жетекшінің қолы: ____________________'+('\nБілім беру ұйымының жауапты тұлғасының қолы: ____________________' if responsible else '')+'\nКүні: ____________________')
def save(d,name):
 d.save(OUT/(name+'.docx'))
 a=Document(OUT/(name+'.docx'));a.core_properties.author='';a.core_properties.last_modified_by=''
 for p in a.paragraphs:
  if p.text.startswith(('1-автордың қолы:','Кестені нақтылап растау:')):
   p._element.getparent().remove(p._element)
  else:
   for r in p.runs:
    old=r.text;new=old.replace(AUTHORS,'1-автор және 2-автор').replace(LEADER,'ғылыми жетекші').replace(SCHOOL,'мектеп')
    if new!=old:r.text=new
 a.save(OUT/(name+'-anonymous.docx'))

# Appendix 3: retain the actual service/report fields blank.
d=create('АНЫҚТАМА\nЖобаның мәтінін алынған материалдарға сәйкестігі (антиплагиат) тексеру нәтижелері туралы','3');common(d)
d.add_paragraph('Қатысушылар санаты:\n□ Бастауыш буын (2–4 сыныптар)\n☑ Орта буын (5–8 сыныптар)')
for s in ['Алынған материалдарды тексеру сервисінің атауы және сілтемесі: __________________________________','Мәтіннің бірегейлігі (%): ____________________','Тексеру күні: ____________________']:d.add_paragraph(s)
d.add_paragraph('Тексеру нәтижесі мен сараптамалық қорытынды: ____________________\nЕсеп нөмірі немесе тіркелген файл: ____________________\nЖауапты тұлға, қызметі: ____________________\nҚолы: ____________________')
d.add_paragraph('Толтыруға арналған бланк. Тексеру өткізілгені мен пайыздық көрсеткіш расталмаған. Нақты сервис есебімен толтырылады.')
save(d,'06-Qalam-antiplagiarism-form')
# Appendix 4: six fields per use case; honest scope of actual assistance.
d=create('Жасанды интеллектіні пайдалану туралы декларация','4');common(d)
d.add_paragraph('ЖИ пайдалану мәртебесі:\n□ Жасанды интеллект пайдаланылған жоқ.\n☑ Жасанды интеллект пайдаланылды.')
questions=['ЖИ құралының атауы және нұсқасы (егер нұсқасы белгілі болса)','Жобамен жұмыс істеу кезеңі және ЖИ пайдалану мақсаты','Негізгі сұраулардың қысқаша мазмұны','ЖИ-ге берілген негізгі сұраулардың қысқаша мазмұны; нәтиже қолданылған бөлік','Алынған нәтиженің дұрыстығы мен нақтылығын тексеру тәсілдері','Автордың (авторлардың) алынған нәтижеге өз бетінше енгізген өзгерістері мен толықтырулары']
cases=[('Сайт коды',['OpenAI ChatGPT/Codex. Нақты модель нұсқасы: __________.','Сайттың елеулі кодын жасау және түзету.','Кітапхана сайты, аккаунт, брондау, пікір, қор және рейтинг қосу.','web/server.js, accounts.js, library-server.js және өзге сайт файлдары. Нақты үлесті жұмыс журналымен көрсету қажет.','Алты автоматты тест өтті; браузер сценарийлері тексерілді. Бұл тексерулер ассистент арқылы орындалды.','Оқушылардың жеке түзетулері туралы расталған жазба жоқ. Нақты әрекет пен дәлел: __________.']),('Құжаттар мен дизайн',['OpenAI ChatGPT/Codex. Нақты модель нұсқасы: __________.','Ғылыми мәтін, қорғау сөзі, нұсқаулық, презентация, безендіру.','Үлгіге бейімдеу, мәтінді жеңілдету, 28 бет және 190 сөздік аннотация дайындау.','28 беттік жұмыстың мәтіні; қосымша құжаттар; 16 слайд; сайттағы таныстыру және логотип нұсқасы.','Бет саны, сөз саны, формат және файлдардың ашылуы тексерілді. Нақты оқушы сауалнамасы жоқ.','Нақты оқушы өңдеуі расталмаған. Өз өзгерісі: __________.'])]
for label,values in cases:
 d.add_paragraph('Қолдану жағдайы: '+label)
 grid(d,['№','Сұралатын мәліметтер','ЖИ пайдаланылуы туралы мәліметтер'],[[i,q,v] for i,(q,v) in enumerate(zip(questions,values),1)])
d.add_paragraph(RULE_NOTE);signatures(d);save(d,'02-Qalam-AI-declaration')
# Appendix 5: all checklist marks await real confirmation.
d=create('Зерттеудің этикасы мен қауіпсіздігі жөніндегі бақылау парағы','5');common(d)
items=['Жобаның ықтимал тәуекелдері алдын ала анықталды','Адамдардың қатысуымен сауалнама/сұхбат жүргізу кезінде ерікті келісім және ата-анасының/заңды өкілінің келісімі алынды','Дербес деректер иесіздендірілді, заңсыз қол жеткізуден қорғалды және ашық ЖИ жүйелеріне енгізілген жоқ','Жануарларға зиян келтіретін әрекеттерге жол берілген жоқ','Химиялық, биологиялық, электрлік және механикалық тәуекелдер болған жағдайда ересек адамның бақылауы қамтамасыз етілді','Пайдаланылған материалдар мен қалдықтарды қауіпсіз кәдеге жарату тәртібі сақталды','Фотосуреттерге, бейнематериалдарға, музыкаға, бейнелерге және бағдарламалық кодқа қатысты авторлық құқықтар сақталды','Зерттеу күнделігі және бастапқы деректер ұсынылды','ЖИ пайдалану туралы декларация толтырылды']
notes=['Қауіп: жеке деректер мен қарапайым бастапқы пароль.','Нақты жазбаша келісім әлі расталмаған.','Тізім бұрын ассистентке берілген. Ашық ЖИ-ге енгізілмеді деген тармақ автоматты расталмайды.','Цифрлық жоба.','Құрылғымен жұмыс тәртібі тексеріледі.','Цифрлық жоба.','Материалдар рұқсаты мен лицензиясы нақтыланады.','Күнделік бланкісі бар; нақты оқушы өлшемдері жоқ.','Декларация жобасы дайын; авторлар тексереді.']
grid(d,['№','Тексеру тармағы','Белгі','Ескерту'],[[i,x,'□ Иә\n□ Қолданылмайды' if i<8 else '□ Иә',n] for i,(x,n) in enumerate(zip(items,notes),1)],10)
signatures(d,True);save(d,'05-Qalam-ethics-safety')
# Appendix 6: actual dates/actions must be written by participants.
d=create('ЗЕРТТЕУ КҮНДЕЛІГІ','6');common(d);d.add_paragraph('Қатысушылар санаты: ☑ Орта буын (5–8 сыныптар)\nЖоба түрі: □ Жеке  ☑ Командалық')
grid(d,['№','Күні немесе кезеңі','Жұмыс кезеңі және мақсаты','Орындалған әрекеттер','Алынған деректер және нәтижелер','Туындаған қиындықтар, өзгерістер және кейінгі әрекеттер','Орындаушы'],[[i,'________',label,'________','________','________','________'] for i,label in enumerate(['Мәселе мен мақсат','Дереккөздер','Каталог пен қор','Брондау және тарих','Сауалнама және сынақ','Талдау және қорытынды','Қорғауға дайындық'],1)],10)
d.add_paragraph('Жазбалар нақты орындалған жұмысқа сүйеніп хронологиялық ретпен толтырылады. Күндер мен оқушылардың әрекеттері расталмағандықтан бос қалдырылды. ЖИ көмегі бөлек көрсетіледі.');save(d,'03-Qalam-research-diary')
# Appendix 7: preserve the seven required stages, never invent student actions.
d=create('КОМАНДАЛЫҚ ЖОБАҒА АВТОРЛЫҚ ҮЛЕС ПАРАҒЫ','7');common(d);d.add_paragraph('Қатысушылар санаты: ☑ Орта буын (5–8 сыныптар)\n1-автор: ____________________  Сынып: 7\n2-автор: ____________________  Сынып: 7')
items=['Тақырыпты таңдау және зерттеу сұрағын тұжырымдау','Дереккөздерді зерделеу','Деректерді жинау немесе эксперимент жүргізу','Нәтижелерді өңдеу және талдау','Өнімді, модельді немесе прототипті жасау (болған жағдайда)','Қорытындыларды тұжырымдау және жұмыс нәтижелерін бағалау','Жұмыс мәтінін, презентацияны және қорғауды дайындау']
grid(d,['№','Жұмыс кезеңі','1-автордың үлесі','2-автордың үлесі','Растайтын материалдар'],[[i,x,'________','________','________'] for i,x in enumerate(items,1)],10)
d.add_paragraph('Нақты үлес код, жазба, өлшем, файл не түсіндіру арқылы расталады. ЖИ жасаған бөліктер дербес оқушы еңбегі ретінде белгіленбейді.');signatures(d);save(d,'08-Qalam-author-contribution')
# Reviewer/supervisor letters are extra draft materials, not official certified forms.
for name in ['04-Qalam-review-draft','07-Qalam-supervisor-opinion-draft']:
 title,content=supplements[name];d=create(title,'Қосымша материал')
 for kind,v in content:
  if kind=='p':d.add_paragraph(v)
  elif kind=='h':d.add_paragraph(v).runs[0].bold=True
  elif kind=='table':grid(d,*v)
 save(d,name)
print('Seven supplementary forms prepared; anonymous copies included.')

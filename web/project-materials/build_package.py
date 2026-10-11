"""Rebuild editable Qalam science materials; render PDFs with LibreOffice afterwards."""
from pathlib import Path
import shutil,math,json
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from pptx import Presentation
from pptx.util import Inches,Pt as PPTPt
from pptx.dml.color import RGBColor as PPTColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from package_content import *
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'downloads'/'Qalam-project-package';OUT.mkdir(parents=True,exist_ok=True)
ASSETS=Path(__file__).parent/'assets';ASSETS.mkdir(exist_ok=True)
for name in ['poster.png','book.png','catalog.png']:
 incoming=Path('/tmp/qalam-project-assets')/name
 if incoming.exists():shutil.copy2(incoming,ASSETS/name)
font_path='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
bold_path='/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
GREEN='#285643';CREAM='#FAFBF3';GOLD='#B8954F';LIGHT='#E8EEDF'
def font(size,bold=False):return ImageFont.truetype(bold_path if bold else font_path,size)
def centered(draw,xy,text,size=25,bold=False,color=GREEN):
 draw.multiline_text(xy,text,font=font(size,bold),fill=color,anchor='mm',align='center',spacing=9)
def diagram(name,labels,subtitle):
 im=Image.new('RGB',(1500,420),CREAM);dr=ImageDraw.Draw(im);dr.rounded_rectangle((12,12,1488,408),24,outline='#CFDDC7',width=3)
 dr.text((38,30),subtitle,font=font(24,True),fill=GREEN)
 n=len(labels);boxw=380 if n==3 else 300;gap=(1420-n*boxw)/(n-1) if n>1 else 0
 for i,(title,body) in enumerate(labels):
  x=40+i*(boxw+gap);dr.rounded_rectangle((x,112,x+boxw,324),20,fill=LIGHT,outline='#AAC1A6',width=2)
  centered(dr,(x+boxw/2,165),title,27,True);centered(dr,(x+boxw/2,250),body,22)
  if i<n-1:
   a=x+boxw+12;b=x+boxw+gap-12;dr.line((a,220,b,220),fill=GOLD,width=5);dr.polygon([(b,220),(b-14,211),(b-14,229)],fill=GOLD)
 im.save(ASSETS/name)
diagram('architecture.png',[('Браузер','Бет пен батырмалар'),('Node.js сервері','Рұқсат және әрекет'),('SQLite','Кітаптар мен тарих')],'QALAM · АҚПАРАТ АҒЫНЫ')
diagram('reservation.png',[('Брондау','3 күнге сақтау'),('Кітап беру','Қор саны азаяды'),('Қайтару','Қор саны артады')],'МЕРЗІМІ АЯҚТАЛСА НЕМЕСЕ ТОҚТАТЫЛСА — БОС ДАНА ҚАЙТА АШЫЛАДЫ')

def page_field(p):
 r=p.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
def setup_doc(landscape=False,body=14):
 doc=Document();s=doc.sections[0];s.page_width=Cm(29.7 if landscape else 21);s.page_height=Cm(21 if landscape else 29.7)
 if landscape:s.orientation=WD_ORIENT.LANDSCAPE
 s.top_margin=Cm(1.6 if landscape else 2);s.bottom_margin=Cm(1.6 if landscape else 2);s.left_margin=Cm(1.6 if landscape else 2.5);s.right_margin=Cm(1.6 if landscape else 2)
 normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(body);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(7)
 for style,size in [('Title',24),('Heading 1',17),('Heading 2',14)]:
  st=doc.styles[style];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('285643');st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(10)
 header=s.header.paragraphs[0];header.text='QALAM · Мектеп кітапханасы';header.style='Caption'
 footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('Qalam · ');page_field(footer)
 doc.core_properties.author=AUTHORS;doc.core_properties.language='kk-KZ'
 return doc

def table(doc,headers,rows,size=11):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,text in zip(t.rows[0].cells,headers):c.text=text
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for c,value in zip(t.add_row().cells,row):c.text=str(value)
 for ri,row in enumerate(t.rows):
  prop=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(prop)
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
    for r in p.runs:r.font.size=Pt(size);r.bold=ri==0
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEDF');cell._tc.get_or_add_tcPr().append(sh)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

def block(doc,item,small=False):
 kind,value=item
 if kind=='p':doc.add_paragraph(value)
 elif kind=='h':doc.add_heading(value,2)
 elif kind=='bullets':
  for text in value:doc.add_paragraph('• '+text)
 elif kind=='table':table(doc,*value,size=10 if small else 11)
 elif kind=='image':
  name,caption,width=value;path=ASSETS/name
  im=Image.open(path);width=min(width,9*im.width/im.height)
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Cm(width))
  cap=doc.add_paragraph(caption);cap.style='Caption';cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

main=setup_doc();main.core_properties.title=TITLE
for i,(title,content) in enumerate(pages):
 if i:main.add_page_break()
 if i==0:
  p=main.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(ROOT/'school-logo.png'),width=Cm(2.3))
  main.add_heading(title,0)
 else:main.add_heading(title,1)
 for item in content:block(main,item)
 if i==0:
  for p in main.paragraphs:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
main.save(OUT/'01-Qalam-scientific-project-25-pages.docx')
for name,(title,content) in supplements.items():
 doc=setup_doc(landscape=name.startswith(('03','08')),body=11 if name.startswith(('03','08')) else 12)
 doc.core_properties.title=title;doc.add_heading(title,0)
 for item in content:block(doc,item,small=name.startswith(('03','05','08')))
 doc.save(OUT/(name+'.docx'))

# Editable widescreen presentation; no commands or fabricated field results.
r=Presentation();r.slide_width=Inches(13.333);r.slide_height=Inches(7.5)
def color(h):return PPTColor.from_string(h.replace('#',''))
def rect(sl,x,y,w,h,fill,line=None,radius=False):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=color(fill)
 if line:sh.line.color.rgb=color(line)
 else:sh.line.fill.background()
 return sh

def text(sl,x,y,w,h,value,size=22,bold=False,fill=GREEN):
 shape=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(value.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=PPTPt(size);p.font.bold=bold;p.font.color.rgb=color(fill);p.space_after=PPTPt(7)
 return shape

def picture(sl,path,x,y,w,h):
 im=Image.open(path);scale=min(w/im.width,h/im.height);pw=im.width*scale;ph=im.height*scale
 sl.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))
for i,(title,subtitle,cards,image) in enumerate(slides,1):
 sl=r.slides.add_slide(r.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=color(CREAM)
 rect(sl,0,0,.15,7.5,GREEN);text(sl,.55,.35,11,.25,'QALAM · ҒЫЛЫМИ ЖОБА · 7-СЫНЫП',11,True)
 sl.shapes.add_picture(str(ROOT/'school-logo.png'),Inches(12),Inches(.25),width=Inches(.65))
 text(sl,.55,.98,12.1,1.1,title,32 if len(title)<50 else 29,True);text(sl,.55,2.03,12,.55,subtitle,20,False,'#718269')
 has_image=image is not None
 leftw=6.4 if has_image else 12.1
 for j,(head,body) in enumerate(cards):
  y=2.85+j*1.13;rect(sl,.55,y,leftw,1.02,LIGHT,radius=True);text(sl,.77,y+.12,leftw-.4,.29,head,18,True);text(sl,.77,y+.47,leftw-.4,.48,body,15 if has_image else 19)
 if has_image:
  rect(sl,7.2,2.85,5.57,3.28,'#FFFFFF','#D4DFCC',True);picture(sl,ASSETS/image,7.3,2.95,5.37,3.08)
  text(sl,7.28,6.22,5.4,.35,'Сайт көрінісі / жүйе сызбасы · оқушы сынағы емес',10,False,'#718269')
 rect(sl,.55,7.0,12.2,.015,'#CEDBC7');text(sl,.55,7.12,11,.22,'Мектеп кітапханасын цифрландыру · '+SCHOOL,9,False,'#718269');text(sl,12,7.1,.7,.25,f'{i:02} / 16',10,True)
 sl.notes_slide.notes_text_frame.text='Сөйлеушіге: '+subtitle+'\n'+'\n'.join(h+': '+b for h,b in cards)+'\nНақты сауалнама көрсеткіші берілмеген. Бағдарламалық тестті оқушы зерттеуімен араластырмаңыз.'
r.core_properties.title=TITLE;r.core_properties.author=AUTHORS;r.save(OUT/'09-Qalam-presentation-16-slides.pptx')
manifest={'main_pages':25,'presentation_slides':16,'authors':AUTHORS,'leader':LEADER,'field_study':'No supplied participant results; forms remain unfilled','automated_tests':'6 executed, 6 passed','signature_documents':'Drafts/forms; signatures and antiplagiarism percentage intentionally unfilled'}
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print('Generated:',len(list(OUT.glob('*.docx'))),'Word documents and 16-slide PowerPoint')

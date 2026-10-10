"""Generate the user guide from USER-GUIDE-KK.md; no application changes."""
from pathlib import Path
from xml.sax.saxutils import escape
import re
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
root=Path(__file__).resolve().parent
source=(root/'USER-GUIDE-KK.md').read_text()
blocks=[]
for group in source.strip().split('\n\n'):
 if group.startswith('# '):blocks.append(('title',group[2:]))
 elif group.startswith('## '):blocks.append(('heading',group[3:]))
 elif group.startswith('- ') or re.match(r'^\d+\. ',group):
  for line in group.splitlines():blocks.append(('step',line))
 else:blocks.append(('body',group.replace('\n',' ')))
doc=Document();doc.core_properties.title='Qalam сайтын қолдану нұсқаулығы';doc.core_properties.language='kk-KZ'
section=doc.sections[0];section.top_margin=Inches(.7);section.bottom_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(8)
for name,size in [('Title',24),('Heading 1',15)]:
 style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('285842')
pdfmetrics.registerFont(TTFont('GuideKazakh','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('GuideKazakhBold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
styles={
 'title':ParagraphStyle('title',fontName='GuideKazakhBold',fontSize=23,leading=30,spaceAfter=18,textColor='#285842',keepWithNext=True),
 'heading':ParagraphStyle('heading',fontName='GuideKazakhBold',fontSize=14,leading=20,spaceBefore=16,spaceAfter=10,textColor='#285842',keepWithNext=True),
 'body':ParagraphStyle('body',fontName='GuideKazakh',fontSize=10.5,leading=16,spaceAfter=10),
 'step':ParagraphStyle('step',fontName='GuideKazakh',fontSize=10.5,leading=16,leftIndent=10,spaceAfter=5)
}
flow=[]
for kind,text in blocks:
 if kind=='title':doc.add_heading(text,0)
 elif kind=='heading':doc.add_heading(text,1)
 else:
  p=doc.add_paragraph(text)
  if kind=='step':p.paragraph_format.left_indent=Inches(.15)
 flow.append(Paragraph(escape(text),styles[kind]))
footer=section.footer.paragraphs[0];footer.text='Qalam · Қолдану нұсқаулығы';footer.style='Caption'
out=root/'downloads';out.mkdir(exist_ok=True);doc.save(out/'Qalam-user-guide.docx')
def page_footer(canvas,document):
 canvas.saveState();canvas.setFont('GuideKazakh',9);canvas.setFillColorRGB(.35,.45,.36)
 canvas.drawString(40,24,'Qalam · Қолдану нұсқаулығы');canvas.drawRightString(A4[0]-40,24,str(document.page));canvas.restoreState()
SimpleDocTemplate(str(out/'Qalam-user-guide.pdf'),pagesize=A4,leftMargin=40,rightMargin=40,topMargin=38,bottomMargin=45,title='Qalam сайтын қолдану нұсқаулығы',author='Qalam').build(flow,onFirstPage=page_footer,onLaterPages=page_footer)
print('Generated Word and PDF:',sum(kind=='heading' for kind,_ in blocks),'sections')

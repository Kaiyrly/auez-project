"""Generate matching speaker script downloads from defense-data.js."""
import json,pathlib,subprocess,copy
from docx import Document
from docx.shared import Pt
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from xml.sax.saxutils import escape
root=pathlib.Path(__file__).resolve().parent
slides=json.loads(subprocess.check_output(['node','-e','process.stdout.write(JSON.stringify(require("./defense-data.js")))'],cwd=root))
assert len(slides)==6 and all(isinstance(s['spoken'],str) for s in slides)
blocks=[('h1','Qalam: қысқартылған қорғау сөзі'),('p','«Онлайн кітапхана: кітапхана жұмысын цифрландыру»'),('p','«Қаныш Сәтпаев атындағы Жәнтікей орта мектебі» КММ'),('p','Авторлар: Мұхаметкәрім Айбын Таңатұлы, Амантай Елжан Ерланұлы. Ғылыми жетекші: Мұхаметкәрім Қаныш Саятұлы.'),('p','6 қысқа бөлім. Мәтін презентациямен бірдей ретпен жазылған.')]
for i,s in enumerate(slides,1):
 blocks.append(('h2',f'{i}-бөлім. {s["title"]} · {s["speaker"]}'))
 blocks.append(('p',s['spoken']))
 if i==4:
  blocks.append(('h1','АЙБЫН ДА, ЕЛЖАН ДА: САЙТТЫҢ ФУНКЦИЯЛАРЫН ӨЗ СӨЗІМЕН ТҮСІНДІРУ'))
  blocks.append(('p','Бұл — екеуіне арналған нұсқау. Осы жерде сайттың не істейтінін өз сөздеріңмен түсіндіріңдер. Содан кейін 5-бөлімге өтіңдер.'))
blocks.extend([('h2','Дайындық нұсқауы — қорғауда оқылмайды'),('p','Презентацияны ашып, бөлімдердің ретімен кезектесіп сөйлеңіздер. Алдын ала дауыстап оқып көріңіздер.')])
md=[];doc=Document();doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(12)
pdfmetrics.registerFont(TTFont('Kazakh','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'));pdfmetrics.registerFont(TTFont('KazakhBold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
styles={'p':ParagraphStyle('p',fontName='Kazakh',fontSize=11,leading=17,spaceAfter=9),'action':ParagraphStyle('action',fontName='Kazakh',fontSize=10,leading=15,textColor='#32665b',spaceAfter=12)}
for k,size in [('h1',18),('h2',14),('h3',12)]:styles[k]=ParagraphStyle(k,fontName='KazakhBold',fontSize=size,leading=size+6,spaceBefore=10,spaceAfter=10,keepWithNext=True)
pdf=[]
for kind,text in blocks:
 md.append(('#'*int(kind[1])+' ' if kind.startswith('h') else '')+text)
 if kind.startswith('h'):doc.add_heading(text,level=int(kind[1]))
 else:doc.add_paragraph(text)
 pdf.append(Paragraph(escape(text),styles[kind]))
(root/'DEFENSE-7-MINUTES.md').write_text('\n\n'.join(md)+'\n')
(root/'downloads').mkdir(exist_ok=True)
for name in ['Qalam-defense-v4','Qalam-defense-v3','Qalam-defense-v2','Qalam-7-minutes']:
 doc.save(root/'downloads'/f'{name}.docx')
 SimpleDocTemplate(str(root/'downloads'/f'{name}.pdf'),leftMargin=40,rightMargin=40,topMargin=36,bottomMargin=36).build(copy.deepcopy(pdf))
print('Generated matching Markdown, Word and PDF')
